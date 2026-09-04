"""Offline safety and capability contract for PDF/Office helpers.

These tests freeze the current behavior of existing library helpers without
adding production parsing, OCR, decryption, macros, formula execution or
format conversion. All fixtures are generated in memory.
"""

import io
import os
import socket
import subprocess
import urllib.request

from io import BytesIO

from docx import Document
from openpyxl import Workbook

from crawler.parser.format import detect_format, parse_content
from crawler.parser.pdf import extract_pdf_text
from crawler.parser.office import extract_office_text

import crawler.parser.pdf as pdf_module
import crawler.parser.office as office_module


# ---------------------------------------------------------------------------
# Minimal in-memory PDF builder
# ---------------------------------------------------------------------------

def _build_pdf(pages):
    """Build a small original PDF with one Helvetica text line per page."""
    count = len(pages)
    font_num = 3 + count
    objects = {}
    objects[1] = "<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join("%d 0 R" % (3 + i) for i in range(count))
    objects[2] = "<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, count)
    for i in range(count):
        page_num = 3 + i
        content_num = font_num + 1 + i
        objects[page_num] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
            % (font_num, content_num)
        )
    objects[font_num] = "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for i, marker in enumerate(pages):
        content_num = font_num + 1 + i
        stream = "BT /F1 12 Tf 72 720 Td (%s) Tj ET" % marker
        objects[content_num] = (
            "<< /Length %d >>\nstream\n%s\nendstream" % (len(stream.encode("ascii")), stream)
        )

    buf = bytearray()
    buf.extend(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = {}
    for num in range(1, max(objects) + 1):
        offsets[num] = len(buf)
        buf.extend(("%d 0 obj\n" % num).encode("ascii"))
        obj = objects[num]
        buf.extend(obj.encode("ascii") if isinstance(obj, str) else obj)
        buf.extend(b"\nendobj\n")
    xref_offset = len(buf)
    buf.extend(("xref\n0 %d\n" % (max(objects) + 1)).encode("ascii"))
    buf.extend(b"0000000000 65535 f \n")
    for num in range(1, max(objects) + 1):
        buf.extend(("%010d 00000 n \n" % offsets[num]).encode("ascii"))
    buf.extend(
        ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (max(objects) + 1, xref_offset)).encode("ascii")
    )
    return bytes(buf)


def _docx_bytes():
    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph("Second paragraph")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "TABLE-TEXT"
    doc.sections[0].header.paragraphs[0].text = "HEADER-TEXT"
    doc.sections[0].footer.paragraphs[0].text = "FOOTER-TEXT"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "SheetA"
    ws.append(["名称", 0, False, None, 1.5])
    ws2 = wb.create_sheet("SheetB")
    ws2.append(["公式单元格", "=SUM(A1:A2)"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Format detection
# ---------------------------------------------------------------------------

def test_detect_pdf_mime():
    assert detect_format("application/pdf") == "pdf"


def test_detect_pdf_mime_with_parameters():
    assert detect_format("application/pdf; charset=utf-8") == "pdf"


def test_detect_mime_case_normalization():
    assert detect_format("APPLICATION/PDF") == "pdf"


def test_detect_doc_docx_mime_routes_office():
    assert detect_format("application/msword") == "office"
    assert detect_format("application/vnd.openxmlformats-officedocument.wordprocessingml.document") == "office"


def test_detect_xls_xlsx_mime_routes_office():
    assert detect_format("application/vnd.ms-excel") == "office"
    assert detect_format("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet") == "office"


def test_detect_ppt_pptx_mime_routes_office():
    assert detect_format("application/vnd.ms-powerpoint") == "office"
    assert detect_format("application/vnd.openxmlformats-officedocument.presentationml.presentation") == "office"


def test_detect_url_extension_fallback():
    assert detect_format("", "https://example.gov.cn/report.pdf") == "pdf"
    assert detect_format("", "https://example.gov.cn/report.docx") == "office"
    assert detect_format("", "https://example.gov.cn/report.xlsx") == "office"


def test_detect_uppercase_extension():
    assert detect_format("", "https://example.gov.cn/report.PDF") == "pdf"


def test_detect_octet_stream_with_pdf_extension():
    assert detect_format("application/octet-stream", "https://example.gov.cn/report.pdf") == "pdf"


def test_detect_mime_wins_over_extension():
    assert detect_format("application/pdf", "https://example.gov.cn/report.docx") == "pdf"


def test_detect_query_fragment_extension_current_limitation():
    # Current implementation does not strip query/fragment before extension lookup.
    assert detect_format("", "https://example.gov.cn/report.pdf?download=1") == "html"


def test_detect_unknown_falls_back_html():
    assert detect_format("", "https://example.gov.cn/noext") == "html"
    assert detect_format("application/x-unknown", "") == "html"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_pdf_single_page_text():
    pdf = _build_pdf(["TASK019C1PDFMARKER"])
    text = extract_pdf_text(pdf)
    assert isinstance(text, str)
    assert text == "TASK019C1PDFMARKER"


def test_pdf_multipage_order():
    pdf = _build_pdf(["TASK019C1PAGEA", "TASK019C1PAGEB"])
    text = extract_pdf_text(pdf)
    assert text == "TASK019C1PAGEA\nTASK019C1PAGEB"


def test_pdf_blank_page_returns_empty_text():
    pdf = _build_pdf([""])
    assert extract_pdf_text(pdf) == ""


def test_pdf_empty_bytes_returns_empty():
    assert extract_pdf_text(b"") == ""


def test_pdf_corrupt_bytes_returns_empty():
    assert extract_pdf_text(b"not a pdf") == ""


def test_pdf_encrypted_or_unsupported_fails_safely():
    # Deliberately not a valid encrypted PDF; contract is safe empty failure.
    raw = b"%PDF-1.4\n1 0 obj\n<< /Encrypt 2 0 R >>\nendobj\n%%EOF"
    assert extract_pdf_text(raw) == ""


def test_pdf_input_bytes_not_mutated():
    pdf = _build_pdf(["TASK019C1PDFMARKER"])
    before = pdf
    extract_pdf_text(pdf)
    assert pdf == before


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def test_docx_paragraphs_in_order():
    text = extract_office_text(_docx_bytes(), "office")
    assert text == "第一段\nSecond paragraph"


def test_docx_table_header_footer_currently_ignored():
    text = extract_office_text(_docx_bytes(), "office")
    assert "TABLE-TEXT" not in text
    assert "HEADER-TEXT" not in text
    assert "FOOTER-TEXT" not in text


def test_docx_empty_document():
    doc = Document()
    buf = BytesIO()
    doc.save(buf)
    assert extract_office_text(buf.getvalue(), "office") == ""


def test_docx_corrupt_bytes_returns_empty():
    assert extract_office_text(b"not a docx", "office") == ""


def test_parse_content_docx_mime_routes_to_office():
    text = parse_content(_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert text == "第一段\nSecond paragraph"


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def test_xlsx_direct_helper_works():
    text = extract_office_text(_xlsx_bytes(), "xlsx")
    assert "名称" in text
    assert "1.5" in text
    assert "公式单元格" in text
    assert "=SUM(A1:A2)" in text


def test_xlsx_sheet_order():
    text = extract_office_text(_xlsx_bytes(), "xlsx")
    assert "SheetA" not in text  # sheet names are not emitted by current helper
    assert text.index("公式单元格") > text.index("名称")


def test_xlsx_zero_false_none_currently_dropped():
    # Current implementation filters falsy cells with `if c`.
    text = extract_office_text(_xlsx_bytes(), "xlsx")
    assert "0" not in text.split(" ")
    assert "False" not in text
    assert "None" not in text


def test_xlsx_formula_kept_as_string():
    text = extract_office_text(_xlsx_bytes(), "xlsx")
    assert "=SUM(A1:A2)" in text
    assert "3" not in text


def test_xlsx_corrupt_bytes_returns_empty():
    assert extract_office_text(b"not an xlsx", "xlsx") == ""


def test_parse_content_xlsx_current_limitation():
    # detect_format routes xlsx MIME to 'office', but unified entry only reaches
    # the XLSX branch when fmt == 'xlsx'; this is a current limitation.
    assert parse_content(_xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet") == ""
    assert parse_content(_xlsx_bytes(), "", "https://example.gov.cn/report.xlsx") == ""


# ---------------------------------------------------------------------------
# Unsupported binary Office formats
# ---------------------------------------------------------------------------

def test_unsupported_binary_office_safe_empty():
    for fmt in ("office", "docx", "xlsx"):
        assert extract_office_text(b"not a real office file", fmt) == ""


def test_detect_unsupported_binary_mimes_only_detects_office():
    for mime in (
        "application/msword",
        "application/vnd.ms-powerpoint",
        "application/vnd.ms-excel",
    ):
        assert detect_format(mime) == "office"


def test_parse_content_unsupported_binary_office_returns_empty():
    assert parse_content(b"legacy binary", "application/msword") == ""
    assert parse_content(b"legacy binary", "application/vnd.ms-powerpoint") == ""
    assert parse_content(b"legacy binary", "application/vnd.ms-excel") == ""


# ---------------------------------------------------------------------------
# Side-effect safety
# ---------------------------------------------------------------------------

def test_parsing_has_no_network_subprocess_or_file_side_effects(monkeypatch):
    def deny(*args, **kwargs):
        raise AssertionError("side effect denied")

    pdf = _build_pdf(["TASK019C1PDFMARKER"])
    docx_data = _docx_bytes()
    xlsx_data = _xlsx_bytes()

    monkeypatch.setattr(socket, "socket", deny)
    monkeypatch.setattr(subprocess, "Popen", deny)
    monkeypatch.setattr(subprocess, "run", deny)
    monkeypatch.setattr(os, "system", deny)
    monkeypatch.setattr(urllib.request, "urlopen", deny)
    monkeypatch.setattr("builtins.open", deny)

    assert extract_pdf_text(pdf) == "TASK019C1PDFMARKER"
    assert extract_office_text(docx_data, "office") == "第一段\nSecond paragraph"
    assert "名称" in extract_office_text(xlsx_data, "xlsx")
    assert parse_content(pdf, "application/pdf") == "TASK019C1PDFMARKER"


def test_no_macro_or_formula_execution_static_source():
    pdf_source = io.open(pdf_module.__file__, encoding="utf-8").read()
    office_source = io.open(office_module.__file__, encoding="utf-8").read()
    for source in (pdf_source, office_source):
        assert "subprocess" not in source
        assert "os.system" not in source
        assert "requests" not in source
        assert "urllib" not in source
        assert "socket" not in source
        assert "open(" not in source
    assert "data_only" not in office_source