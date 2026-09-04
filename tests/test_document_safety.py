"""Tests for crawler.parser.document_safety safety gate."""

import dataclasses
import io
import os
import socket
import subprocess
import urllib.request
import zipfile

from io import BytesIO

import pytest

from docx import Document
from openpyxl import Workbook

from crawler.parser import document_safety as ds
from crawler.parser.document_safety import (
    DEFAULT_DOCUMENT_SAFETY_POLICY,
    DocumentParseResult,
    DocumentSafetyPolicy,
    safe_parse_document,
)

from tests.test_document_format_contract import _build_pdf, _docx_bytes, _xlsx_bytes


def _zip_bytes(entries):
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries:
            zf.writestr(name, data)
    return buf.getvalue()


def _docx_many_paragraphs(count):
    doc = Document()
    for i in range(count):
        doc.add_paragraph("段落%d" % i)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_two_sheets():
    wb = Workbook()
    ws = wb.active
    ws.append(["a"])
    wb.create_sheet("Second")
    ws2 = wb["Second"]
    ws2.append(["b"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_many_rows(count):
    wb = Workbook()
    ws = wb.active
    for i in range(count):
        ws.append([i])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_many_cells(count):
    wb = Workbook()
    ws = wb.active
    ws.append(list(range(count)))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _encrypted_zip_bytes():
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        info = zipfile.ZipInfo("[Content_Types].xml")
        info.flag_bits |= 0x1
        zf.writestr(info, b"x")
        info2 = zipfile.ZipInfo("word/document.xml")
        info2.flag_bits |= 0x1
        zf.writestr(info2, b"x")
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Policy
# ---------------------------------------------------------------------------

def test_default_policy_values():
    p = DEFAULT_DOCUMENT_SAFETY_POLICY
    assert p.max_input_bytes == 20 * 1024 * 1024
    assert p.max_pdf_pages == 500
    assert p.max_archive_members == 2048
    assert p.max_archive_uncompressed_bytes == 100 * 1024 * 1024
    assert p.max_archive_member_bytes == 50 * 1024 * 1024
    assert p.max_compression_ratio == 100
    assert p.max_docx_paragraphs == 20_000
    assert p.max_xlsx_sheets == 32
    assert p.max_xlsx_rows == 100_000
    assert p.max_xlsx_cells == 1_000_000
    assert p.max_output_chars == 2_000_000


def test_policy_immutable():
    p = DEFAULT_DOCUMENT_SAFETY_POLICY
    with pytest.raises(dataclasses.FrozenInstanceError):
        p.max_input_bytes = 1  # type: ignore[misc]


@pytest.mark.parametrize("value", [0, -1, 1.5, "10", True])
def test_policy_rejects_invalid_values(value):
    with pytest.raises(ValueError):
        DocumentSafetyPolicy(max_input_bytes=value)


def test_custom_small_policy():
    p = DocumentSafetyPolicy(max_input_bytes=10, max_pdf_pages=1, max_output_chars=100)
    assert p.max_input_bytes == 10


# ---------------------------------------------------------------------------
# Type, ordering, format
# ---------------------------------------------------------------------------

def test_non_bytes_rejected():
    result = safe_parse_document("text")  # type: ignore[arg-type]
    assert result.status == "rejected"
    assert result.reason == "invalid_input_type"
    assert result.content == ""


def test_empty_bytes():
    result = safe_parse_document(b"")
    assert result.status == "empty"
    assert result.reason == "empty_input"


def test_input_too_large_before_unsupported():
    policy = DocumentSafetyPolicy(max_input_bytes=5)
    result = safe_parse_document(b"123456", "application/msword", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "input_too_large"


def test_unknown_format_unsupported():
    result = safe_parse_document(b"x", "")
    assert result.status == "unsupported"
    assert result.reason == "unsupported_format"
    assert result.extraction_method == "none"


def test_mime_priority_over_extension():
    result = safe_parse_document(_build_pdf(["MARKER"]), "application/pdf", "https://example.gov.cn/x.docx")
    assert result.status == "success"
    assert result.document_format == "pdf"


def test_query_fragment_extension_limitation_preserved():
    result = safe_parse_document(_build_pdf(["MARKER"]), "", "https://example.gov.cn/report.pdf?x=1")
    assert result.status == "unsupported"
    assert result.reason == "unsupported_format"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_pdf_success():
    result = safe_parse_document(_build_pdf(["TASK019C2PDF"]), "application/pdf")
    assert result.status == "success"
    assert result.reason == "ok"
    assert result.content == "TASK019C2PDF"
    assert result.extraction_method == "pdf"


def test_pdf_blank_empty():
    result = safe_parse_document(_build_pdf([""]), "application/pdf")
    assert result.status == "empty"
    assert result.reason == "empty_content"
    assert result.content == ""


def test_pdf_corrupt_failed():
    result = safe_parse_document(b"%PDF-corrupt", "application/pdf")
    assert result.status == "rejected"
    assert result.reason == "parse_failed"


def test_pdf_signature_mismatch():
    result = safe_parse_document(b"not pdf", "application/pdf")
    assert result.status == "rejected"
    assert result.reason == "signature_mismatch"


def test_pdf_page_limit_equal_allowed():
    policy = DocumentSafetyPolicy(max_pdf_pages=1)
    result = safe_parse_document(_build_pdf(["PAGE"]), "application/pdf", policy=policy)
    assert result.status == "success"


def test_pdf_page_limit_exceeded(monkeypatch):
    policy = DocumentSafetyPolicy(max_pdf_pages=1)
    def deny(*args, **kwargs):
        raise AssertionError("extractor called")
    monkeypatch.setattr(ds, "extract_pdf_text", deny)
    result = safe_parse_document(_build_pdf(["PAGE1", "PAGE2"]), "application/pdf", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "pdf_page_limit"
    assert result.content == ""


def test_pdf_output_too_large():
    policy = DocumentSafetyPolicy(max_output_chars=5)
    result = safe_parse_document(_build_pdf(["TASK019C2PDF"]), "application/pdf", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "output_too_large"
    assert result.content == ""


# ---------------------------------------------------------------------------
# ZIP preflight
# ---------------------------------------------------------------------------

def test_non_zip_signature_mismatch():
    result = safe_parse_document(b"not a zip", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.status == "rejected"
    assert result.reason == "signature_mismatch"


def test_zip_missing_core_member():
    raw = _zip_bytes([("[Content_Types].xml", b"x")])
    result = safe_parse_document(raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.status == "rejected"
    assert result.reason == "invalid_archive"


def test_zip_member_count_limit():
    raw = _zip_bytes([("[Content_Types].xml", b"a"), ("word/document.xml", b"b"), ("extra", b"c")])
    policy = DocumentSafetyPolicy(max_archive_members=2)
    result = safe_parse_document(raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "archive_member_limit"


def test_zip_member_size_limit():
    raw = _zip_bytes([("[Content_Types].xml", b"a"), ("word/document.xml", b"b" * 20)])
    policy = DocumentSafetyPolicy(max_archive_member_bytes=10)
    result = safe_parse_document(raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "archive_member_size_limit"


def test_zip_uncompressed_limit():
    raw = _zip_bytes([("[Content_Types].xml", b"a"), ("word/document.xml", b"b" * 20)])
    policy = DocumentSafetyPolicy(max_archive_uncompressed_bytes=10)
    result = safe_parse_document(raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "archive_uncompressed_limit"


def test_zip_compression_ratio_limit():
    raw = _zip_bytes([("[Content_Types].xml", b"a"), ("word/document.xml", b"b" * 1000)])
    policy = DocumentSafetyPolicy(max_compression_ratio=2)
    result = safe_parse_document(raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "compression_ratio_limit"


@pytest.mark.parametrize("name", ["../evil", "a/../../evil", "C:/evil", "/evil", "a\\..\\evil"])
def test_zip_unsafe_paths(name):
    raw = _zip_bytes([("[Content_Types].xml", b"a"), ("word/document.xml", b"b"), (name, b"c")])
    result = safe_parse_document(raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.status == "rejected"
    assert result.reason == "unsafe_archive_path"


def test_unsafe_path_nul_rejected_by_helper():
    from crawler.parser.document_safety import _unsafe_archive_path
    assert _unsafe_archive_path("bad\x00name") is True


def test_zip_encrypted_member(monkeypatch):
    class FakeInfo:
        def __init__(self, name):
            self.filename = name
            self.flag_bits = 0x1
            self.file_size = 1
            self.compress_size = 1
            self.external_attr = 0

    class FakeZip:
        def __init__(self, *args, **kwargs):
            pass
        def infolist(self):
            return [FakeInfo("[Content_Types].xml"), FakeInfo("word/document.xml")]
        def close(self):
            pass

    raw = _encrypted_zip_bytes()
    monkeypatch.setattr(ds.zipfile, "is_zipfile", lambda *a, **k: True)
    monkeypatch.setattr(ds.zipfile, "ZipFile", FakeZip)
    result = safe_parse_document(raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.status == "rejected"
    assert result.reason == "encrypted_archive_member"


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def test_docx_success():
    result = safe_parse_document(_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.status == "success"
    assert result.content == "第一段\nSecond paragraph"
    assert result.extraction_method == "docx"


def test_docx_empty_document():
    doc = Document()
    buf = BytesIO()
    doc.save(buf)
    result = safe_parse_document(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.status == "empty"
    assert result.reason == "empty_content"


def test_docx_paragraph_limit(monkeypatch):
    policy = DocumentSafetyPolicy(max_docx_paragraphs=2)
    def deny(*args, **kwargs):
        raise AssertionError("extractor called")
    monkeypatch.setattr(ds, "extract_office_text", deny)
    result = safe_parse_document(_docx_many_paragraphs(3), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "docx_paragraph_limit"


def test_docx_output_too_large():
    policy = DocumentSafetyPolicy(max_output_chars=3)
    result = safe_parse_document(_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "output_too_large"


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def test_xlsx_success_and_formula_not_executed():
    result = safe_parse_document(_xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert result.status == "success"
    assert "=SUM(A1:A2)" in result.content
    assert result.extraction_method == "xlsx"


def test_xlsx_sheet_limit(monkeypatch):
    policy = DocumentSafetyPolicy(max_xlsx_sheets=1)
    def deny(*args, **kwargs):
        raise AssertionError("extractor called")
    monkeypatch.setattr(ds, "extract_office_text", deny)
    result = safe_parse_document(_xlsx_two_sheets(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "xlsx_sheet_limit"


def test_xlsx_row_limit(monkeypatch):
    policy = DocumentSafetyPolicy(max_xlsx_rows=2)
    def deny(*args, **kwargs):
        raise AssertionError("extractor called")
    monkeypatch.setattr(ds, "extract_office_text", deny)
    result = safe_parse_document(_xlsx_many_rows(3), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "xlsx_row_limit"


def test_xlsx_cell_limit(monkeypatch):
    policy = DocumentSafetyPolicy(max_xlsx_cells=2)
    def deny(*args, **kwargs):
        raise AssertionError("extractor called")
    monkeypatch.setattr(ds, "extract_office_text", deny)
    result = safe_parse_document(_xlsx_many_cells(3), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "xlsx_cell_limit"


def test_xlsx_output_too_large():
    policy = DocumentSafetyPolicy(max_output_chars=3)
    result = safe_parse_document(_xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", policy=policy)
    assert result.status == "rejected"
    assert result.reason == "output_too_large"


# ---------------------------------------------------------------------------
# Result and side effects
# ---------------------------------------------------------------------------

def test_unsupported_format_not_parsed(monkeypatch):
    def deny(*args, **kwargs):
        raise AssertionError("extractor called")
    monkeypatch.setattr(ds, "extract_office_text", deny)
    result = safe_parse_document(b"x", "application/msword")
    assert result.status == "unsupported"
    assert result.reason == "unsupported_format"


def test_rejected_path_does_not_call_extractor(monkeypatch):
    def deny(*args, **kwargs):
        raise AssertionError("extractor called")
    monkeypatch.setattr(ds, "extract_office_text", deny)
    raw = _zip_bytes([("../../evil", b"x")])
    result = safe_parse_document(raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.status == "rejected"
    assert result.reason == "unsafe_archive_path"


def test_all_failures_have_empty_content():
    results = [
        safe_parse_document("x"),  # type: ignore[arg-type]
        safe_parse_document(b""),
        safe_parse_document(b"123456", "application/msword", DocumentSafetyPolicy(max_input_bytes=5)),
        safe_parse_document(b"x", "application/msword"),
    ]
    assert all(r.content == "" for r in results)


def test_result_deterministic():
    raw = _build_pdf(["DETERMINISTIC"])
    a = safe_parse_document(raw, "application/pdf")
    b = safe_parse_document(raw, "application/pdf")
    assert a == b


def test_no_network_subprocess_or_file_side_effects(monkeypatch):
    pdf = _build_pdf(["SIDE"])
    docx = _docx_bytes()
    xlsx = _xlsx_bytes()

    def deny(*args, **kwargs):
        raise AssertionError("side effect denied")

    monkeypatch.setattr(socket, "socket", deny)
    monkeypatch.setattr(subprocess, "Popen", deny)
    monkeypatch.setattr(subprocess, "run", deny)
    monkeypatch.setattr(os, "system", deny)
    monkeypatch.setattr(urllib.request, "urlopen", deny)
    monkeypatch.setattr("builtins.open", deny)

    assert safe_parse_document(pdf, "application/pdf").status == "success"
    assert safe_parse_document(docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document").status == "success"
    assert safe_parse_document(xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet").status == "success"


def test_no_zip_extract_in_source():
    source = io.open(ds.__file__, encoding="utf-8").read()
    assert "extract(" not in source
    assert "extractall" not in source