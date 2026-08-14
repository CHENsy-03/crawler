"""Independent safety gate for existing PDF/DOCX/XLSX helpers.

This module is intentionally not imported by any production module. It wraps
the existing extraction helpers with deterministic resource preflight and a
structured result contract.
"""

import re
import zipfile
from dataclasses import dataclass, fields
from io import BytesIO

from docx import Document
from openpyxl import load_workbook
from PyPDF2 import PdfReader

from crawler.parser.pdf import extract_pdf_text
from crawler.parser.office import extract_office_text

DEFAULT_DOCUMENT_SAFETY_POLICY = None


@dataclass(frozen=True)
class DocumentSafetyPolicy:
    max_input_bytes: int = 20 * 1024 * 1024
    max_pdf_pages: int = 500
    max_archive_members: int = 2048
    max_archive_uncompressed_bytes: int = 100 * 1024 * 1024
    max_archive_member_bytes: int = 50 * 1024 * 1024
    max_compression_ratio: int = 100
    max_docx_paragraphs: int = 20_000
    max_xlsx_sheets: int = 32
    max_xlsx_rows: int = 100_000
    max_xlsx_cells: int = 1_000_000
    max_output_chars: int = 2_000_000

    def __post_init__(self) -> None:
        for field in fields(self):
            value = getattr(self, field.name)
            if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
                raise ValueError("DocumentSafetyPolicy %s must be a positive integer" % field.name)


DEFAULT_DOCUMENT_SAFETY_POLICY = DocumentSafetyPolicy()


@dataclass(frozen=True)
class DocumentParseResult:
    status: str = ""
    reason: str = ""
    document_format: str = ""
    extraction_method: str = "none"
    content: str = ""
    input_bytes: int = 0
    output_chars: int = 0
    pdf_pages: int = 0
    archive_members: int = 0
    archive_uncompressed_bytes: int = 0
    docx_paragraphs: int = 0
    xlsx_sheets: int = 0
    xlsx_rows: int = 0
    xlsx_cells: int = 0


class _SafetyRejection(Exception):
    def __init__(self, reason: str, **stats: int) -> None:
        super().__init__(reason)
        self.reason = reason
        self.stats = stats


_SUPPORTED_MIME = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
}

_UNSUPPORTED_MIME = {
    "application/msword": "doc",
    "application/vnd.ms-powerpoint": "ppt",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.ms-excel": "xls",
}

_EXT_FORMATS = {
    "pdf": "pdf",
    "docx": "docx",
    "xlsx": "xlsx",
    "doc": "doc",
    "xls": "xls",
    "ppt": "ppt",
    "pptx": "pptx",
    "docm": "docm",
    "xlsm": "xlsm",
    "pptm": "pptm",
}

_DOCX_REQUIRED = {"[Content_Types].xml", "word/document.xml"}
_XLSX_REQUIRED = {"[Content_Types].xml", "xl/workbook.xml"}


def _detect_format(content_type: str, url: str) -> str:
    mime = ""
    if content_type:
        mime = content_type.split(";", 1)[0].strip().lower()
    if mime:
        if mime in _SUPPORTED_MIME:
            return _SUPPORTED_MIME[mime]
        if mime in _UNSUPPORTED_MIME:
            return "unsupported"
    ext = url.rsplit(".", 1)[-1].lower() if "." in url else ""
    return _EXT_FORMATS.get(ext, "unknown")


def _result(status: str, reason: str, document_format: str = "", extraction_method: str = "none",
            content: str = "", stats: dict | None = None) -> DocumentParseResult:
    data = dict(stats or {})
    return DocumentParseResult(
        status=status,
        reason=reason,
        document_format=document_format,
        extraction_method=extraction_method,
        content=content,
        input_bytes=int(data.get("input_bytes", 0)),
        output_chars=len(content),
        pdf_pages=int(data.get("pdf_pages", 0)),
        archive_members=int(data.get("archive_members", 0)),
        archive_uncompressed_bytes=int(data.get("archive_uncompressed_bytes", 0)),
        docx_paragraphs=int(data.get("docx_paragraphs", 0)),
        xlsx_sheets=int(data.get("xlsx_sheets", 0)),
        xlsx_rows=int(data.get("xlsx_rows", 0)),
        xlsx_cells=int(data.get("xlsx_cells", 0)),
    )


def _reject(reason: str, **stats: int) -> _SafetyRejection:
    return _SafetyRejection(reason, **stats)


def _preflight_zip(content: bytes, policy: DocumentSafetyPolicy, required: set[str]) -> dict:
    if not zipfile.is_zipfile(BytesIO(content)):
        raise _reject("signature_mismatch", input_bytes=len(content))
    try:
        archive = zipfile.ZipFile(BytesIO(content))
    except zipfile.BadZipFile:
        raise _reject("invalid_archive", input_bytes=len(content))

    names = set()
    members = 0
    total_uncompressed = 0
    try:
        for info in archive.infolist():
            members += 1
            if members > policy.max_archive_members:
                raise _reject("archive_member_limit", archive_members=members)

            if info.flag_bits & 0x1:
                raise _reject("encrypted_archive_member", archive_members=members)

            name = info.filename or ""
            if _unsafe_archive_path(name):
                raise _reject("unsafe_archive_path", archive_members=members)

            is_symlink = (info.external_attr >> 16) & 0o170000 == 0o120000
            if is_symlink:
                raise _reject("unsafe_archive_path", archive_members=members)

            if info.file_size > policy.max_archive_member_bytes:
                raise _reject("archive_member_size_limit", archive_members=members)

            if info.file_size > 0:
                ratio = info.file_size / max(info.compress_size, 1)
                if ratio > policy.max_compression_ratio:
                    raise _reject("compression_ratio_limit", archive_members=members)

            total_uncompressed += info.file_size
            if total_uncompressed > policy.max_archive_uncompressed_bytes:
                raise _reject("archive_uncompressed_limit", archive_uncompressed_bytes=total_uncompressed)

            normalized = name.replace("\\", "/")
            names.add(normalized)
    finally:
        archive.close()

    if not required.issubset(names):
        raise _reject("invalid_archive", archive_members=members)

    return {
        "archive_members": members,
        "archive_uncompressed_bytes": total_uncompressed,
        "input_bytes": len(content),
    }


def _unsafe_archive_path(name: str) -> bool:
    if not name or "\x00" in name:
        return True
    if name.startswith("/") or re.match(r"^[A-Za-z]:", name):
        return True
    normalized = name.replace("\\", "/")
    parts = [part for part in normalized.split("/") if part not in ("", ".")]
    return ".." in parts


def _parse_pdf(content: bytes, policy: DocumentSafetyPolicy) -> DocumentParseResult:
    if not content.startswith(b"%PDF-"):
        raise _reject("signature_mismatch", input_bytes=len(content))

    stats = {"input_bytes": len(content)}
    try:
        reader = PdfReader(BytesIO(content))
    except Exception:
        raise _reject("parse_failed", **stats)

    pages = len(reader.pages)
    stats["pdf_pages"] = pages
    if reader.is_encrypted:
        raise _reject("encrypted_document", **stats)
    if pages > policy.max_pdf_pages:
        raise _reject("pdf_page_limit", **stats)

    text = extract_pdf_text(content)
    if not text:
        return _result("empty", "empty_content", "pdf", "pdf", "", stats)
    if len(text) > policy.max_output_chars:
        raise _reject("output_too_large", **stats)
    return _result("success", "ok", "pdf", "pdf", text, stats)


def _parse_docx(content: bytes, policy: DocumentSafetyPolicy) -> DocumentParseResult:
    stats = _preflight_zip(content, policy, _DOCX_REQUIRED)
    try:
        doc = Document(BytesIO(content))
    except Exception:
        raise _reject("invalid_archive", **stats)

    paragraphs = len(doc.paragraphs)
    stats["docx_paragraphs"] = paragraphs
    if paragraphs > policy.max_docx_paragraphs:
        raise _reject("docx_paragraph_limit", **stats)

    text = extract_office_text(content, "docx")
    if not text:
        return _result("empty", "empty_content", "docx", "docx", "", stats)
    if len(text) > policy.max_output_chars:
        raise _reject("output_too_large", **stats)
    return _result("success", "ok", "docx", "docx", text, stats)


def _parse_xlsx(content: bytes, policy: DocumentSafetyPolicy) -> DocumentParseResult:
    stats = _preflight_zip(content, policy, _XLSX_REQUIRED)
    wb = None
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
        sheets = len(wb.worksheets)
        stats["xlsx_sheets"] = sheets
        if sheets > policy.max_xlsx_sheets:
            raise _reject("xlsx_sheet_limit", **stats)

        rows = 0
        cells = 0
        for ws in wb.worksheets:
            max_row = int(ws.max_row or 0)
            max_col = int(ws.max_column or 0)
            if max_row > policy.max_xlsx_rows:
                raise _reject("xlsx_row_limit", **stats)
            if max_col > 0 and max_row > 0 and max_row * max_col > policy.max_xlsx_cells:
                raise _reject("xlsx_cell_limit", **stats)
            for row in ws.iter_rows(values_only=True):
                rows += 1
                if rows > policy.max_xlsx_rows:
                    raise _reject("xlsx_row_limit", **stats)
                cells += len(row)
                if cells > policy.max_xlsx_cells:
                    raise _reject("xlsx_cell_limit", **stats)

        stats["xlsx_rows"] = rows
        stats["xlsx_cells"] = cells
    finally:
        if wb is not None:
            wb.close()

    text = extract_office_text(content, "xlsx")
    if not text:
        return _result("empty", "empty_content", "xlsx", "xlsx", "", stats)
    if len(text) > policy.max_output_chars:
        raise _reject("output_too_large", **stats)
    return _result("success", "ok", "xlsx", "xlsx", text, stats)


def safe_parse_document(
    content: bytes,
    content_type: str = "",
    url: str = "",
    policy: DocumentSafetyPolicy = DEFAULT_DOCUMENT_SAFETY_POLICY,
) -> DocumentParseResult:
    if not isinstance(policy, DocumentSafetyPolicy):
        raise ValueError("policy must be a DocumentSafetyPolicy")
    # Dataclass __post_init__ already validates values; explicit check for config misuse.
    if not isinstance(content, bytes):
        return _result("rejected", "invalid_input_type", "unknown", "none", "")
    if not content:
        return _result("empty", "empty_input", "unknown", "none", "", {"input_bytes": 0})
    if len(content) > policy.max_input_bytes:
        return _result("rejected", "input_too_large", "unknown", "none", "", {"input_bytes": len(content)})

    document_format = _detect_format(content_type, url)
    if document_format not in ("pdf", "docx", "xlsx"):
        return _result("unsupported", "unsupported_format", document_format, "none", "", {"input_bytes": len(content)})

    try:
        if document_format == "pdf":
            return _parse_pdf(content, policy)
        if document_format == "docx":
            return _parse_docx(content, policy)
        return _parse_xlsx(content, policy)
    except _SafetyRejection as exc:
        return _result("rejected", exc.reason, document_format, "none", "", {"input_bytes": len(content), **exc.stats})
    except Exception:
        return _result("failed", "parse_failed", document_format, "none", "", {"input_bytes": len(content)})